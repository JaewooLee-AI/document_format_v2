"""DOCX(OOXML) 읽기/쓰기 어댑터. HwpxDocument와 동일한 (table_index, row, col) 좌표 인터페이스를 제공한다.

python-docx 의 docDefaults(styles.xml)가 이미 "나눔휴먼OTF Regular / 9pt"로 설정되어 있어
(SVI_forms.md 본문·입력 내용 규격과 일치), 내용 셀에 새 run을 추가할 때 별도 rPr 없이도
올바른 서식이 적용된다. 라벨 셀은 건드리지 않는다.
"""
import copy
import io

from docx import Document
from docx.oxml.ns import qn


class DocxDocument:
    def __init__(self, document: Document):
        self._doc = document
        self._tables = document.tables

    @classmethod
    def open(cls, data: bytes) -> "DocxDocument":
        return cls(Document(io.BytesIO(data)))

    def table_count(self) -> int:
        return len(self._tables)

    def _cell(self, table_index, row, col):
        return self._tables[table_index].cell(row, col)

    def get_cell_text(self, table_index, row, col) -> str:
        return self._cell(table_index, row, col).text

    def set_cell_text(self, table_index, row, col, text: str):
        cell = self._cell(table_index, row, col)
        cell.text = ""
        lines = text.split("\n") if text else [""]
        first = True
        for line in lines:
            if first:
                cell.paragraphs[0].add_run(line)
                first = False
            else:
                cell.add_paragraph(line)

    # -- 문단/런 접근 (위촉장 전용) -----------------------------------------
    def get_paragraph_run_text(self, table_index, row, col, paragraph_index, run_index) -> str:
        cell = self._cell(table_index, row, col)
        return cell.paragraphs[paragraph_index].runs[run_index].text or ""

    def set_paragraph_run_text(self, table_index, row, col, paragraph_index, run_index, text: str):
        cell = self._cell(table_index, row, col)
        cell.paragraphs[paragraph_index].runs[run_index].text = text

    # -- 반복행(표) 크기 조정 -----------------------------------------------
    def ensure_row_count(self, table_index, header_row: int, desired_data_rows: int):
        tbl = self._tables[table_index]._tbl
        trs = tbl.findall(qn("w:tr"))
        data_trs = trs[header_row + 1:]
        current = len(data_trs)
        if desired_data_rows > current:
            template_tr = data_trs[-1] if data_trs else trs[header_row]
            for _ in range(current, desired_data_rows):
                new_tr = copy.deepcopy(template_tr)
                tbl.append(new_tr)
        elif desired_data_rows < current:
            for tr in data_trs[desired_data_rows:]:
                tbl.remove(tr)
        self._tables = self._doc.tables

    def get_row_count(self, table_index, header_row: int) -> int:
        tbl = self._tables[table_index]._tbl
        trs = tbl.findall(qn("w:tr"))
        return max(0, len(trs) - header_row - 1)

    # -- 라벨/값 행이 번갈아 나오는 표(회의록 위원의견 등) 쌍 개수 조정 -------
    def set_pair_count(self, table_index, desired_pairs: int):
        tbl = self._tables[table_index]._tbl
        trs = tbl.findall(qn("w:tr"))
        current_pairs = len(trs) // 2
        if desired_pairs > current_pairs:
            label_tpl, value_tpl = trs[0], trs[1]
            for _ in range(current_pairs, desired_pairs):
                tbl.append(copy.deepcopy(label_tpl))
                tbl.append(copy.deepcopy(value_tpl))
        elif desired_pairs < current_pairs:
            for tr in trs[desired_pairs * 2:]:
                tbl.remove(tr)
        self._tables = self._doc.tables

    def get_pair_count(self, table_index) -> int:
        tbl = self._tables[table_index]._tbl
        return len(tbl.findall(qn("w:tr"))) // 2

    # -- 사진 삽입 ------------------------------------------------------
    def insert_image(self, table_index, row, col, image_stream, height_cm: float):
        from docx.shared import Cm
        cell = self._cell(table_index, row, col)
        cell.text = ""
        run = cell.paragraphs[0].add_run()
        run.add_picture(image_stream, height=Cm(height_cm))

    # -- 저장 -----------------------------------------------------------
    def to_bytes(self) -> bytes:
        buf = io.BytesIO()
        self._doc.save(buf)
        return buf.getvalue()
